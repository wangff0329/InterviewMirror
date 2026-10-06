import hashlib
import io
import json
import re
import sqlite3
from contextlib import closing
from datetime import datetime

from docx import Document
from langchain_openai import ChatOpenAI

from app.config import Settings
from app.schemas import ModelConfig


def init_profile_database(settings: Settings) -> None:
    with closing(sqlite3.connect(settings.database_path)) as connection:
        with connection:
            connection.execute(
                """
                CREATE TABLE IF NOT EXISTS profile_documents (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    user_id INTEGER NOT NULL,
                    filename TEXT NOT NULL,
                    content_type TEXT NOT NULL,
                    file_hash TEXT NOT NULL,
                    extracted_text TEXT NOT NULL,
                    created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
                )
                """
            )
            connection.execute(
                """
                CREATE TABLE IF NOT EXISTS profile_versions (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    user_id INTEGER NOT NULL,
                    document_id INTEGER NOT NULL,
                    version_number INTEGER NOT NULL,
                    profile_json TEXT NOT NULL,
                    created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
                    UNIQUE(user_id, version_number)
                )
                """
            )


def extract_document_text(content: bytes, filename: str, content_type: str) -> str:
    suffix = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    if content_type == "application/pdf" or suffix == "pdf":
        from app.services.pdf import extract_pdf_text

        return extract_pdf_text(content)
    if (
        content_type
        == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        or suffix == "docx"
    ):
        document = Document(io.BytesIO(content))
        paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
        table_text = [
            " | ".join(cell.text.strip() for cell in row.cells)
            for table in document.tables
            for row in table.rows
            if any(cell.text.strip() for cell in row.cells)
        ]
        text = "\n".join(paragraphs + table_text)
        if not text:
            raise ValueError("Word 文档中没有可提取的文本。")
        return text[:50_000]
    raise ValueError("只支持 PDF 或 DOCX 文件。")


def extract_profile(text: str) -> dict:
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    joined = "\n".join(lines)
    email_match = re.search(r"[\w.+-]+@[\w-]+\.[\w.-]+", joined)
    phone_match = re.search(r"(?:\+?86[-\s]?)?1[3-9]\d{9}", joined)
    age_match = re.search(r"(?:年龄|age)\s*[:：]?\s*(\d{2})", joined, re.IGNORECASE)
    target_match = re.search(r"(?:求职意向|目标岗位|target role)\s*[:：]?\s*(.+)", joined, re.IGNORECASE)
    name = ""
    if lines and not re.search(r"简历|resume|curriculum", lines[0], re.IGNORECASE):
        name = re.split(r"[|｜,\s]", lines[0])[0][:40]

    sections = _split_resume_sections(lines)
    education = _extract_education(
        sections.get("education", []) or sections.get("education_work", []),
        lines,
    )
    work_experience = _extract_experience(sections.get("work", []), "工作经历")
    project_experience = _extract_experience(sections.get("project", []), "项目经历")
    skills = _extract_skills(sections.get("skills", []))
    return {
        "basic_info": {
            "name": name,
            "email": email_match.group(0) if email_match else "",
            "phone": phone_match.group(0) if phone_match else "",
            "age": int(age_match.group(1)) if age_match else None,
            "location": "",
        },
        "education": education[:10],
        "work_experience": work_experience,
        "project_experience": project_experience,
        "skills": list(dict.fromkeys(skills))[:30],
        "target_role": target_match.group(1).strip()[:200] if target_match else "",
        "source_text_preview": text[:1000],
        "raw_sections": sections,
        "extraction_note": "当前版本按简历章节保留教育、工作、项目和技能原文；未找到的信息不会自动猜测，请确认后保存。",
    }


def _section_kind(title: str) -> str | None:
    normalized = re.sub(r"[\s:：·•|｜]", "", title).lower()
    if re.search(r"项目|project|作品|portfolio", normalized):
        return "project"
    if re.search(r"教育|学历|education|academic", normalized):
        return "education_work" if re.search(r"工作", normalized) else "education"
    if re.search(r"工作|实习|employment|professional|workexperience", normalized):
        return "work"
    if re.search(r"技能|技术|skill|technology|competenc", normalized):
        return "skills"
    return None


def _split_resume_sections(lines: list[str]) -> dict[str, list[str]]:
    sections = {"education": [], "education_work": [], "work": [], "project": [], "skills": [], "other": []}
    current = "other"
    heading_pattern = re.compile(
        r"^(?:\d+[.)、]\s*)?(教育与工作经历|教育经历|教育背景|工作经历|工作经验|实习经历|"
        r"科研/项目经历|项目经历|项目经验|专业技能|技能|技术栈|"
        r"education|work experience|professional experience|project experience|projects|skills|技术能力)"
        r"(?:（[^）]*）|\([^)]*\))?\s*[:：]?\s*$",
        re.IGNORECASE,
    )
    for line in lines:
        match = heading_pattern.match(line)
        if match:
            current = _section_kind(match.group(1)) or "other"
            continue
        sections[current].append(line)
    return sections


def _extract_education(section_lines: list[str], all_lines: list[str]) -> list[dict[str, str]]:
    source = section_lines or [
        line
        for line in all_lines
        if re.search(
            r"本科|硕士|博士|大专|学士|研究生|大学|university|college|bachelor|master|phd",
            line,
            re.IGNORECASE,
        )
    ]
    return [
        {"school": line[:160], "degree": "", "major": "", "period": ""}
        for line in source[:10]
        if len(line) >= 2
    ]


def _extract_experience(section_lines: list[str], default_title: str) -> list[dict[str, object]]:
    if not section_lines:
        return []
    entries = []
    current = []
    for line in section_lines:
        starts_entry = (
            current
            and re.search(r"20\d{2}[./年-]\d{1,2}.*(?:20\d{2}|至今|present)", line, re.IGNORECASE)
            and not line.startswith(("", "•", "·", "-", "*"))
        )
        if starts_entry:
            entries.append(current)
            current = []
        current.append(line)
    if current:
        entries.append(current)
    return [
        {
            "title": entry[0][:160] if entry else default_title,
            "description": "\n".join(entry[1:])[:3000],
            "source": "\n".join(entry)[:3200],
        }
        for entry in entries[:20]
    ]


def _extract_skills(section_lines: list[str]) -> list[str]:
    skills = []
    for line in section_lines:
        value = re.split(r"[:：]", line, maxsplit=1)[-1]
        skills.extend(item.strip() for item in re.split(r"[,，、|｜/／]", value) if item.strip())
    return skills


def extract_profile_with_ai(text: str, config: ModelConfig, rule_profile: dict) -> dict:
    if config.provider == "demo":
        return rule_profile
    if not config.api_key or not config.base_url:
        raise ValueError("AI 资料解析需要同时配置 API Key 和 Base URL。")
    llm = ChatOpenAI(
        api_key=config.api_key,
        base_url=config.base_url.rstrip("/"),
        model=config.model,
        temperature=0,
    )
    prompt = f"""
你是简历结构化解析器。只根据简历原文提取事实，不得猜测或补全。
必须保留所有教育、工作、项目、论文、奖项和技能信息，尤其不能把项目经历丢到技能里。
只返回 JSON，不要 Markdown，字段必须符合：
{{
  "basic_info": {{"name": "", "email": "", "phone": "", "age": null, "location": ""}},
  "education": [{{"school": "", "degree": "", "major": "", "period": "", "description": ""}}],
  "work_experience": [{{"title": "", "company": "", "period": "", "description": "", "achievements": []}}],
  "project_experience": [{{"title": "", "role": "", "period": "", "description": "", "achievements": [], "technologies": []}}],
  "skills": [],
  "awards": [],
  "academic_results": [],
  "target_role": ""
}}
缺失字段使用空字符串、空数组或 null；年龄只有原文明确出现时才填写。
简历原文：
{text[:50000]}
"""
    response = llm.invoke(prompt)
    parsed = json.loads(response.content.strip().replace("```json", "").replace("```", "").strip())
    parsed["source_text_preview"] = text[:1000]
    parsed["raw_sections"] = rule_profile.get("raw_sections", {})
    parsed["extraction_note"] = "AI 已根据原文完成结构化提取，请核对项目、时间和成果后保存。"
    return parsed


def create_profile_draft(
    settings: Settings,
    user_id: int,
    filename: str,
    content_type: str,
    content: bytes,
) -> dict:
    text = extract_document_text(content, filename, content_type)
    profile = extract_profile(text)
    with closing(sqlite3.connect(settings.database_path)) as connection:
        with connection:
            cursor = connection.execute(
                """
                INSERT INTO profile_documents
                    (user_id, filename, content_type, file_hash, extracted_text)
                VALUES (?, ?, ?, ?, ?)
                """,
                (user_id, filename, content_type, hashlib.sha256(content).hexdigest(), text),
            )
            return {
                "document_id": cursor.lastrowid,
                "filename": filename,
                "profile": profile,
                "required_fields": required_profile_fields(profile),
            }


def required_profile_fields(profile: dict) -> list[dict[str, str]]:
    basic = profile.get("basic_info", {})
    education = profile.get("education") or []
    has_education = any(
        isinstance(item, dict)
        and any(str(item.get(key) or "").strip() for key in ("school", "degree", "major"))
        for item in education
    )
    return [
        {"key": "basic_info.name", "label": "姓名", "value": str(basic.get("name") or "")},
        {"key": "education", "label": "至少一条教育经历", "value": "yes" if has_education else ""},
        {"key": "target_role", "label": "目标岗位", "value": str(profile.get("target_role") or "")},
    ]


def save_profile_version(settings: Settings, user_id: int, document_id: int, profile: dict) -> dict:
    missing = [item["label"] for item in required_profile_fields(profile) if not item["value"].strip()]
    if missing:
        raise ValueError(f"请先补充必填信息：{'、'.join(missing)}")
    with closing(sqlite3.connect(settings.database_path)) as connection:
        with connection:
            document = connection.execute(
                "SELECT id FROM profile_documents WHERE id = ? AND user_id = ?",
                (document_id, user_id),
            ).fetchone()
            if not document:
                raise ValueError("资料文件不存在或不属于当前用户。")
            latest = connection.execute(
                "SELECT COALESCE(MAX(version_number), 0) FROM profile_versions WHERE user_id = ?",
                (user_id,),
            ).fetchone()[0]
            cursor = connection.execute(
                """
                INSERT INTO profile_versions
                    (user_id, document_id, version_number, profile_json)
                VALUES (?, ?, ?, ?)
                """,
                (user_id, document_id, latest + 1, json.dumps(profile, ensure_ascii=False)),
            )
            return {
                "id": cursor.lastrowid,
                "version_number": latest + 1,
                "profile": profile,
                "created_at": datetime.now().isoformat(timespec="seconds"),
            }


def list_profile_versions(settings: Settings, user_id: int) -> list[dict]:
    with closing(sqlite3.connect(settings.database_path)) as connection:
        rows = connection.execute(
            """
            SELECT v.id, v.version_number, v.profile_json, v.created_at, d.filename
            FROM profile_versions v
            JOIN profile_documents d ON d.id = v.document_id
            WHERE v.user_id = ?
            ORDER BY v.version_number DESC
            """,
            (user_id,),
        ).fetchall()
    return [
        {
            "id": row[0],
            "version_number": row[1],
            "profile": json.loads(row[2]),
            "created_at": row[3],
            "filename": row[4],
        }
        for row in rows
    ]


def get_latest_profile_context(settings: Settings, user_id: int) -> dict | None:
    with closing(sqlite3.connect(settings.database_path)) as connection:
        row = connection.execute(
            """
            SELECT v.id, v.version_number, v.profile_json, v.created_at,
                   d.filename, d.extracted_text
            FROM profile_versions v
            JOIN profile_documents d ON d.id = v.document_id
            WHERE v.user_id = ?
            ORDER BY v.version_number DESC
            LIMIT 1
            """,
            (user_id,),
        ).fetchone()
    if not row:
        return None
    return {
        "id": row[0],
        "version_number": row[1],
        "profile": json.loads(row[2]),
        "created_at": row[3],
        "filename": row[4],
        "resume_text": row[5],
    }
