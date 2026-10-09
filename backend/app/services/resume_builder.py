import io
from datetime import datetime

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from app.schemas import ApplicationAdaptation


def build_resume_docx(
    original_text: str,
    adaptation: ApplicationAdaptation,
    source_filename: str,
) -> tuple[bytes, str]:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    normal = document.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal.font.size = Pt(10.5)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run("岗位定制简历")
    title_run.bold = True
    title_run.font.size = Pt(20)

    meta = document.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.add_run(f"目标岗位：{adaptation.target_role}    ").bold = True
    meta.add_run(f"生成时间：{datetime.now():%Y-%m-%d}")

    _add_heading(document, "个人简介")
    document.add_paragraph(adaptation.profile_summary)

    _add_heading(document, "岗位匹配重点")
    document.add_paragraph(adaptation.positioning)

    _add_heading(document, "针对岗位优化的经历要点")
    for bullet in adaptation.tailored_bullets:
        document.add_paragraph(bullet, style="List Bullet")

    _add_heading(document, "原始简历内容（事实核对区）")
    document.add_paragraph(
        f"来源文件：{source_filename}\n以下内容来自原始简历，用于核对生成结果，请根据需要调整排版。"
    )
    for paragraph in original_text.splitlines():
        text = paragraph.strip()
        if text:
            document.add_paragraph(text)

    _add_heading(document, "提交前检查")
    for caution in adaptation.cautions:
        document.add_paragraph(caution, style="List Bullet")

    output = io.BytesIO()
    document.save(output)
    return output.getvalue(), "岗位定制简历.docx"


def _add_heading(document: Document, text: str) -> None:
    heading = document.add_paragraph()
    heading.paragraph_format.space_before = Pt(12)
    heading.paragraph_format.space_after = Pt(4)
    run = heading.add_run(text)
    run.bold = True
    run.font.size = Pt(13)
